"""Exporta a apresentação para PDF pesquisável e PowerPoint visual.

Uso: python scripts/export.py [--out qa/exports]

Abre site/ em modo estático (?still), captura cada slide em 1920×1080 no estado
final das animações e gera:

- site/downloads/agentes-autonomos-hermes-python-brasil-2026.pdf: uma página
  por slide, com a imagem do slide e uma camada de texto invisível para busca;
- site/downloads/agentes-autonomos-hermes-python-brasil-2026.pptx: uma imagem
  por slide, com título, texto do slide e links nas notas.

As imagens embutidas usam JPEG de qualidade 90; o PDF é comparado com a captura
PNG original para garantir que a composição não mudou.

Requer playwright (Chromium), pillow, pymupdf, python-pptx e opencv-python-headless.
"""
from __future__ import annotations

import argparse
import contextlib
import functools
import hashlib
import http.server
import json
import socketserver
import threading
from pathlib import Path
from zipfile import ZipFile

import cv2
import numpy as np
import pymupdf
from build import write_manifest
from PIL import Image, ImageChops, ImageDraw, ImageStat
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / "site"
NAME = "agentes-autonomos-hermes-python-brasil-2026"
TITLE = "Agentes Autônomos de IA com Hermes Agent"
URL = "https://pycodebr.com.br/python-brasil-2026/"
REPOSITORY = "https://github.com/pycodebr/pycodebr_python_brasil_2026"
QR_EXPECTED = {
    "pycodebr": "https://www.instagram.com/pycodebr/",
    "acompanhe": "https://pycodebr.com.br/python-brasil-2026",
    "materiais": REPOSITORY,
}
HIDE_CHROME = "#toolbar,#overview,#menu,#blank{display:none!important}"
SLIDE_TEXT = """
() => {
  const slide = document.querySelector('.slide.active');
  const clone = slide.cloneNode(true);
  clone.querySelectorAll('[aria-hidden="true"],.ghost,script,style,video').forEach(n => n.remove());
  const text = (clone.innerText || clone.textContent || '').replace(/\\s+/g, ' ').trim();
  return { id: slide.dataset.id, title: slide.dataset.title, chapter: slide.dataset.chapter || '', text };
}
"""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, format, *args):
        return


@contextlib.contextmanager
def serve(directory: Path):
    handler = functools.partial(QuietHandler, directory=str(directory))
    with socketserver.TCPServer(("127.0.0.1", 0), handler) as httpd:
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        try:
            yield f"http://127.0.0.1:{httpd.server_address[1]}/"
        finally:
            httpd.shutdown()


def capture(captures: Path) -> list[dict]:
    slides: list[dict] = []
    with serve(SITE) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080}, reduced_motion="reduce")
        errors: list[str] = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)
        page.goto(base + "?still#1")
        page.wait_for_function("document.documentElement.dataset.ready === 'true'")
        page.add_style_tag(content=HIDE_CHROME)
        count = page.evaluate("presentation.count")
        require(not page.evaluate("document.body.classList.contains('reading')"), "Exportação em modo leitura")
        for i in range(count):
            page.evaluate("i => presentation.show(i)", i)
            page.wait_for_timeout(900)
            page.evaluate("document.fonts.ready")
            broken = page.evaluate(
                "[...document.querySelectorAll('.slide.active img')].filter(i=>!i.complete||!i.naturalWidth).map(i=>i.getAttribute('src'))"
            )
            require(not broken, f"Imagem quebrada no slide {i + 1}: {broken}")
            path = captures / f"slide-{i + 1:02d}.png"
            page.screenshot(path=str(path))
            require(Image.open(path).size == (1920, 1080), "Dimensão de captura incorreta")
            Image.open(path).convert("RGB").save(path.with_suffix(".jpg"), quality=90, optimize=True)
            info = page.evaluate(SLIDE_TEXT)
            require(info["title"] and info["text"], f"Slide {i + 1} sem título ou texto")
            info["number"] = i + 1
            slides.append(info)
        require(not errors, f"Erros de execução na exportação: {errors[:3]}")
        browser.close()
    return slides


def decode_qr(path: Path) -> list[str]:
    image = cv2.imdecode(np.frombuffer(path.read_bytes(), np.uint8), cv2.IMREAD_COLOR)
    ok, decoded, _, _ = cv2.QRCodeDetector().detectAndDecodeMulti(image)
    return [d for d in decoded if d] if ok else []


def build_pdf(slides: list[dict], captures: Path, target: Path) -> list[dict]:
    width, height = 20 * 72, 11.25 * 72
    document = pymupdf.open()
    document.set_metadata({"title": TITLE, "author": "Felipe Azambuja · PycodeBR",
                           "subject": "Python Brasil 2026", "keywords": "Hermes Agent, agentes autônomos, Python"})
    for item in slides:
        page = document.new_page(width=width, height=height)
        text = f"{item['title']}. {item['text']} @pycodebr · {URL}"
        box = pymupdf.Rect(36, 36, width - 36, height - 36)
        for size in (11, 9, 7, 5, 4, 3):
            if page.insert_textbox(box, text, fontsize=size, render_mode=3) >= 0:
                break
        else:
            raise RuntimeError(f"Texto do slide {item['number']} não coube na camada pesquisável")
        page.insert_image(page.rect, filename=str(captures / f"slide-{item['number']:02d}.jpg"))
    document.save(target, garbage=4, deflate=True)
    document.close()
    pdf = pymupdf.open(target)
    require(len(pdf) == len(slides), "Número de páginas do PDF divergente")
    comparisons = []
    for item, page in zip(slides, pdf):
        require("@pycodebr" in page.get_text() and item["title"].split()[0] in page.get_text(),
                f"Texto pesquisável ausente na página {item['number']}")
        require(abs(page.rect.width / page.rect.height - 16 / 9) < 0.001, "Proporção do PDF incorreta")
        pix = page.get_pixmap(matrix=pymupdf.Matrix(4 / 3, 4 / 3))
        rendered = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        original = Image.open(captures / f"slide-{item['number']:02d}.png").convert("RGB")
        if rendered.size != original.size:
            rendered = rendered.resize(original.size)
        error = sum(ImageStat.Stat(ImageChops.difference(rendered, original)).mean) / 3
        require(error < 5, f"PDF difere do slide {item['number']}: {error}")
        comparisons.append({"slide": item["number"], "mean_absolute_error": round(error, 4)})
    pdf.close()
    return comparisons


def build_pptx(slides: list[dict], captures: Path, target: Path) -> None:
    deck = Presentation()
    deck.slide_width = Inches(20)
    deck.slide_height = Inches(11.25)
    deck.core_properties.title = TITLE
    deck.core_properties.author = "Felipe Azambuja · PycodeBR"
    deck.core_properties.subject = "Python Brasil 2026"
    deck.core_properties.comments = "Versão estática em imagens por slide. Animações e interações estão na apresentação HTML."
    for item in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_picture(str(captures / f"slide-{item['number']:02d}.jpg"), 0, 0,
                                         deck.slide_width, deck.slide_height)
        shape.name = item["title"]
        shape._element.nvPicPr.cNvPr.set("descr", item["text"][:900])
        slide.notes_slide.notes_text_frame.text = "\n\n".join([
            f"{item['number']}. {item['title']}", item["text"],
            "Apresentação interativa: " + URL, "Repositório: " + REPOSITORY,
        ])
    deck.save(target)
    require(len(Presentation(target).slides) == len(slides), "Número de slides do PPTX divergente")
    with ZipFile(target) as archive:
        require(archive.testzip() is None, "PPTX inválido")


def contact_sheets(captures: Path, output: Path, count: int) -> None:
    for start in range(0, count, 6):
        paths = [captures / f"slide-{n + 1:02d}.png" for n in range(start, min(start + 6, count))]
        sheet = Image.new("RGB", (1280, 390 * ((len(paths) + 1) // 2)), "#03060c")
        draw = ImageDraw.Draw(sheet)
        for index, path in enumerate(paths):
            image = Image.open(path).convert("RGB")
            image.thumbnail((640, 360))
            x, y = index % 2 * 640, index // 2 * 390
            sheet.paste(image, (x, y))
            draw.text((x + 12, y + 365), path.stem, fill="white")
        sheet.save(output / f"contact-{start // 6 + 1}.jpg", quality=90)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(ROOT / "qa" / "exports"))
    args = parser.parse_args()
    output = Path(args.out)
    captures = output / "captures"
    captures.mkdir(parents=True, exist_ok=True)
    downloads = SITE / "downloads"
    downloads.mkdir(exist_ok=True)
    slides = capture(captures)
    qr = {}
    for slide_id, expected in QR_EXPECTED.items():
        item = next(s for s in slides if s["id"] == slide_id)
        decoded = decode_qr(captures / f"slide-{item['number']:02d}.png")
        require(expected in decoded, f"QR Code divergente em {slide_id}: {decoded}")
        qr[slide_id] = expected
    pdf_path, pptx_path = downloads / (NAME + ".pdf"), downloads / (NAME + ".pptx")
    comparisons = build_pdf(slides, captures, pdf_path)
    build_pptx(slides, captures, pptx_path)
    contact_sheets(captures, output, len(slides))
    report = {
        "slides": len(slides), "pdf_pages": len(slides), "pdf_searchable": True,
        "pdf_visual_comparisons": comparisons, "pptx_slides": len(slides),
        "pptx_mode": "imagens por slide, texto e links nas notas", "qr_codes_decoded": qr,
        "files": [{"name": f.name, "bytes": f.stat().st_size, "sha256": hashlib.sha256(f.read_bytes()).hexdigest()}
                  for f in (pdf_path, pptx_path)],
        "passed": True,
    }
    write_manifest()
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: report[k] for k in ("slides", "pdf_pages", "pptx_slides", "qr_codes_decoded", "files", "passed")},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
